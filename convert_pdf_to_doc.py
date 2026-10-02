import fitz
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os
import shutil

def convert_pdf_to_docx():
    pdf_path = 'Shiksha_Setu_CEP_Project_Report.pdf'
    docx_path = 'Shiksha_Setu_Project_Report_v2.docx'

    if not os.path.exists(pdf_path):
        print(f"Error: {pdf_path} not found!")
        return

    doc_pdf = fitz.open(pdf_path)
    print(f"📖 Opened {pdf_path} with {len(doc_pdf)} pages for DOC conversion.")

    doc_word = docx.Document()
    
    # Configure 1 inch margins
    for section in doc_word.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    tmp_img_dir = 'tmp_pdf_extracted_imgs'
    os.makedirs(tmp_img_dir, exist_ok=True)

    for i, page in enumerate(doc_pdf):
        page_num = i + 1
        
        # Add Page Header indicator for DOC alignment
        p_hdr = doc_word.add_paragraph()
        p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_hdr = p_hdr.add_run(f"--- Page {page_num} of {len(doc_pdf)} ---")
        r_hdr.font.name = 'Arial'
        r_hdr.font.size = Pt(8)
        r_hdr.font.color.rgb = RGBColor(148, 163, 184)

        # Extract text blocks
        text_blocks = page.get_text('blocks')
        for b in text_blocks:
            text = b[4].strip()
            if not text:
                continue

            p = doc_word.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15

            # Formatting heuristics
            if text.startswith("Chapter ") or text.startswith("Appendix ") or text == "Abstract" or text == "DECLARATION" or text == "ACKNOWLEDGEMENT" or text == "Plagiarism Reports" or text == "TABLE OF CONTENTS":
                p.paragraph_format.space_before = Pt(14)
                p.paragraph_format.keep_with_next = True
                r = p.add_run(text)
                r.font.name = 'Arial'
                r.font.size = Pt(14)
                r.bold = True
                r.font.color.rgb = RGBColor(5, 150, 105)
            elif text.startswith("1.") or text.startswith("2.") or text.startswith("3.") or text.startswith("4.") or text.startswith("5.") or text.startswith("6.") or text.startswith("7."):
                if len(text) < 90 and not text.endswith('.'):
                    p.paragraph_format.space_before = Pt(10)
                    p.paragraph_format.keep_with_next = True
                    r = p.add_run(text)
                    r.font.name = 'Arial'
                    r.font.size = Pt(12)
                    r.bold = True
                    r.font.color.rgb = RGBColor(13, 148, 136)
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    r = p.add_run(text)
                    r.font.name = 'Calibri'
                    r.font.size = Pt(11)
                    r.font.color.rgb = RGBColor(30, 41, 59)
            elif text.startswith("Figure "):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(text)
                r.font.name = 'Arial'
                r.font.size = Pt(9.5)
                r.bold = True
                r.font.color.rgb = RGBColor(4, 120, 87)
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                r = p.add_run(text)
                r.font.name = 'Calibri'
                r.font.size = Pt(11)
                r.font.color.rgb = RGBColor(30, 41, 59)

        # Extract embedded images on page
        image_list = page.get_images(full=True)
        for img_index, img_info in enumerate(image_list):
            xref = img_info[0]
            base_image = doc_pdf.extract_image(xref)
            image_bytes = base_image["image"]
            image_ext = base_image["ext"]
            img_name = f"extracted_p{page_num}_{img_index}.{image_ext}"
            img_path = os.path.join(tmp_img_dir, img_name)

            with open(img_path, "wb") as f_img:
                f_img.write(image_bytes)

            p_img = doc_word.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(8)
            p_img.paragraph_format.space_after = Pt(8)
            p_img.add_run().add_picture(img_path, width=Inches(5.8))

        if page_num < len(doc_pdf):
            doc_word.add_page_break()

    doc_word.save(docx_path)
    shutil.copy(docx_path, 'Shiksha_Setu_CEP_Project_Report.docx')
    shutil.copy(docx_path, 'public/Shiksha_Setu_CEP_Project_Report.docx')
    shutil.copy(docx_path, 'public/Shiksha_Setu_Project_Report_v2.docx')

    print(f"🎉 Successfully converted {pdf_path} to {docx_path}!")

if __name__ == '__main__':
    convert_pdf_to_docx()
